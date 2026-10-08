"""Build a daily camera report from the same ordered photos as the gallery."""
import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageOps


def _photo_table(doc):
    table = doc.add_table(rows=0, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column in table.columns:
        column.width = Cm(3.6)
    margins = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        margin = OxmlElement(f'w:{side}')
        margin.set(qn('w:w'), '40' if side in ('left', 'right') else '0')
        margin.set(qn('w:type'), 'dxa')
        margins.append(margin)
    table._tbl.tblPr.append(margins)
    return table


def export_daily_photo_word(project_name, day, photos):
    """Return an in-memory DOCX; photos contain path, shift and send_time."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.5)
    section.left_margin = section.right_margin = Cm(1.5)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.keep_with_next = False
    doc.add_paragraph('Tổng hợp ảnh camera theo ngày', 'Title')
    doc.add_paragraph(f'Dự án: {project_name}')
    doc.add_paragraph(f'Ngày: {day.strftime("%d/%m/%Y")} — Tổng số: {len(photos)} ảnh')

    groups = (
        ('Ca sáng · 05:00–15:00', {'morning'}),
        ('Ca chiều · 16:00–24:00', {'afternoon'}),
        ('Ảnh khác', {'outside', 'unknown'}),
    )
    previous_has_photos = False
    for label, shifts in groups:
        selected = [p for p in photos if p['shift'] in shifts]
        if label == 'Ảnh khác' and not selected:
            continue
        heading = doc.add_paragraph(f'{label} ({len(selected)} ảnh)', 'Heading 1')
        heading.paragraph_format.page_break_before = previous_has_photos
        if not selected:
            doc.add_paragraph('Không có ảnh')
            previous_has_photos = False
            continue
        table = _photo_table(doc)
        for offset in range(0, len(selected), 5):
            # Four bounded-height rows fit inside A4 with the report heading.
            # Start a fresh table rather than letting Word choose print breaks.
            if offset and offset % 20 == 0:
                continuation = doc.add_paragraph(f'{label} (tiếp theo)', 'Heading 1')
                continuation.paragraph_format.page_break_before = True
                doc.add_paragraph(f'{project_name} — {day.strftime("%d/%m/%Y")}')
                table = _photo_table(doc)
            row = table.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell, photo in zip(row.cells, selected[offset:offset + 5]):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # Do not chain table rows together: Word otherwise moves the
                # image grid away from the report header on the first page.
                paragraph.paragraph_format.keep_with_next = False
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                try:
                    with Image.open(photo['path']) as original:
                        image = ImageOps.exif_transpose(original)
                        image.load()
                        # Normalize unsupported formats and keep camera reports compact.
                        image = image.convert('RGB')
                        image.thumbnail((1200, 1600))
                        stream = io.BytesIO()
                        image.save(stream, format='JPEG', quality=90, optimize=True)
                        stream.seek(0)
                        scale = min(3.45 / image.width, 4.5 / image.height)
                        paragraph.add_run().add_picture(
                            stream, width=Cm(image.width * scale), height=Cm(image.height * scale))
                except (OSError, ValueError, Image.DecompressionBombError) as exc:
                    raise ValueError(f'Không thể xuất ảnh "{photo["path"].name}": tệp hỏng hoặc không đọc được') from exc
        previous_has_photos = True

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', project_name).strip().rstrip('.')
    return output, f'Anh_{safe_name}_{day.isoformat()}.docx'
