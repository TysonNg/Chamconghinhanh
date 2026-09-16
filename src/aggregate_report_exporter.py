# -*- coding: utf-8 -*-
"""Tạo dữ liệu và file Word giải trình chấm công tổng hợp."""

import os
import re
import unicodedata
from datetime import date, datetime
from typing import Dict, Iterable, List, Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image


_DATE_FORMATS = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y")
_ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_TEMPLATE_PATH = os.path.join(
    _ROOT_DIR, "templates", "word", "giai_trinh_tong_hop.docx"
)
DATA_HEADERS = ["TÊN", "NGÀY", "GIẢI TRÌNH", "HÌNH ẢNH THỰC TẾ", "GHI CHÚ"]
DATA_WIDTHS = [Cm(3.2), Cm(2.2), Cm(4.4), Cm(5.0), Cm(3.4)]


def _parse_date(value) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _explanation_for(record: Dict) -> str:
    if record.get("same_in_out"):
        return "Giờ vào và giờ ra trùng hoặc quá gần, bổ sung hình ảnh thực tế"
    if record.get("under_3h"):
        return "Làm dưới 3 tiếng, bổ sung hình ảnh thực tế"
    if record.get("missing_checkin"):
        return "Thiếu giờ vào, bổ sung hình ảnh thực tế"
    if record.get("missing_checkout"):
        return "Thiếu giờ ra, bổ sung hình ảnh thực tế"
    if record.get("is_absent"):
        return "Vắng mặt, cần giải trình"
    return "Cần giải trình dữ liệu chấm công"


def _note_for(record: Dict) -> str:
    parts = []
    if record.get("gio_vao"):
        parts.append(f"Giờ vào: {record['gio_vao']}")
    if record.get("gio_ra"):
        parts.append(f"Giờ ra: {record['gio_ra']}")
    if record.get("review_reason"):
        parts.append(record["review_reason"])
    elif not record.get("matched_image_path"):
        parts.append("Không tìm thấy ảnh phù hợp")
    return "; ".join(parts)


def collect_issue_rows(
    persons: Iterable[Dict],
    from_date=None,
    to_date=None,
) -> List[Dict]:
    """Gom, lọc và sắp xếp các ngày chấm công cần giải trình."""
    start = _parse_date(from_date)
    end = _parse_date(to_date)
    rows = []
    for person in persons:
        for record in person.get("records", []):
            is_issue = any(
                record.get(key)
                for key in ("is_absent", "missing_checkin", "missing_checkout", "same_in_out", "under_3h", "record_conflict")
            )
            if not is_issue:
                continue
            record_date = _parse_date(record.get("date"))
            if record_date is None:
                continue
            if start and record_date < start:
                continue
            if end and record_date > end:
                continue
            rows.append({
                "name": str(person.get("name") or "").strip(),
                "date": record_date.strftime("%d/%m/%Y"),
                "date_value": record_date,
                "explanation": _explanation_for(record),
                "note": _note_for(record),
                "image_path": record.get("matched_image_path"),
            })
    rows.sort(key=lambda item: (item["name"].casefold(), item["date_value"]))
    return rows


def _set_font(run, size=10, bold=None, italic=None, color=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def _set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def _set_table_layout_fixed(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_row_repeat(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def _add_page_field(paragraph, field_name):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_name
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char, instr, separate, text, end))
    _set_font(run, size=9)


def build_default_template(template_path=DEFAULT_TEMPLATE_PATH):
    """Tạo template A4 dọc có thể chỉnh sửa độc lập với mã xuất báo cáo."""
    os.makedirs(os.path.dirname(template_path), exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("GIẢI TRÌNH CHẤM CÔNG NHÂN VIÊN")
    _set_font(run, size=15, bold=True, color=(0, 0, 0))

    period = doc.add_paragraph()
    period.alignment = WD_ALIGN_PARAGRAPH.CENTER
    period.paragraph_format.space_after = Pt(3)
    _set_font(
        period.add_run("Từ ngày {{FROM_DATE}} đến ngày {{TO_DATE}}"),
        size=12,
        bold=True,
    )

    subject = doc.add_paragraph()
    subject.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subject.paragraph_format.space_after = Pt(8)
    _set_font(
        subject.add_run(
            "(V/v: nhân viên bảo vệ Dự án: {{PROJECT_NAME}} thiếu dữ liệu chấm công vân tay)"
        ),
        size=10.5,
        italic=True,
    )

    intro_lines = [
        "Kính gửi: Ban Quản lý khu vực Dự án: {{PROJECT_NAME}}",
        (
            "Công ty TNHH DV Bảo vệ Thế An thực hiện yêu cầu của Ban Quản lý khu vực "
            "Dự án: {{PROJECT_NAME}} về việc nhân viên bấm dấu vân tay khi vào làm việc "
            "và khi ra về. Qua đối chiếu bảng công thực tế, một số trường hợp còn thiếu "
            "dữ liệu chấm công. Công ty xin giải trình như sau:"
        ),
    ]
    for index, text in enumerate(intro_lines):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Cm(0.75) if index else Cm(0)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.15
        _set_font(paragraph.add_run(text), size=10.5, bold=(index == 0))

    table = doc.add_table(rows=2, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_layout_fixed(table)
    _set_table_borders(table)
    for index, (cell, text, width) in enumerate(zip(table.rows[0].cells, DATA_HEADERS, DATA_WIDTHS)):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_margins(cell, top=100, bottom=100)
        _set_cell_shading(cell, "1F4E78")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        _set_font(paragraph.add_run(text), size=8.5, bold=True, color=(255, 255, 255))
    _set_row_repeat(table.rows[0])
    marker = table.rows[1]
    marker.cells[0].merge(marker.cells[-1])
    marker.cells[0].text = "{{DATA_ROWS}}"

    closing_texts = [
        (
            "Công ty TNHH DV Bảo vệ Thế An cam kết nhân viên đơn vị bảo vệ đi làm đầy đủ "
            "trong các thời gian nêu trên. Kính mong Ban Quản lý và các phòng ban xem xét."
        ),
        "Trân trọng cảm ơn!",
        "Ý kiến khác: ............................................................................................................................",
    ]
    for index, text in enumerate(closing_texts):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.space_before = Pt(5 if index == 0 else 2)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_with_next = True
        _set_font(paragraph.add_run(text), size=10.5, italic=(index == 2))

    signature = doc.add_table(rows=4, cols=2)
    signature.alignment = WD_TABLE_ALIGNMENT.CENTER
    signature.autofit = False
    _remove_table_borders(signature)
    date_cell = signature.rows[0].cells[0].merge(signature.rows[0].cells[1])
    date_p = date_cell.paragraphs[0]
    date_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(date_p.add_run("{{REPORT_DATE_LONG}}"), size=10.5, italic=True)
    roles = ["XÁC NHẬN CỦA BAN QUẢN LÝ", "ĐẠI DIỆN CÔNG TY BẢO VỆ THẾ AN"]
    names = ["TRẦN HOÀNG LÊ KHIẾT", "NGUYỄN TRÚC PHƯƠNG"]
    for index in range(2):
        signature.rows[1].cells[index].width = Cm(9.1)
        role_p = signature.rows[1].cells[index].paragraphs[0]
        role_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        role_p.paragraph_format.keep_with_next = True
        _set_font(role_p.add_run(roles[index]), size=10, bold=True)
        blank_p = signature.rows[2].cells[index].paragraphs[0]
        blank_p.paragraph_format.space_after = Pt(58)
        name_p = signature.rows[3].cells[index].paragraphs[0]
        name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(name_p.add_run(names[index]), size=10, bold=True)
    for row in signature.rows:
        _set_row_cant_split(row)
    for row in signature.rows[:3]:
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.keep_with_next = True

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(footer_p.add_run("Trang "), size=9)
    _add_page_field(footer_p, "PAGE")
    _set_font(footer_p.add_run(" / "), size=9)
    _add_page_field(footer_p, "NUMPAGES")

    doc.save(template_path)
    return template_path


def _replace_tokens(doc, replacements):
    def replace_in_paragraph(paragraph):
        for run in paragraph.runs:
            for token, value in replacements.items():
                if token in run.text:
                    run.text = run.text.replace(token, value)

    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph)


def _find_data_table(doc):
    for table in doc.tables:
        if table.rows and [cell.text.strip() for cell in table.rows[0].cells] == DATA_HEADERS:
            return table
    raise ValueError("Template không có bảng dữ liệu giải trình hợp lệ")


def _fit_image(image_path, max_width_cm=4.6, max_height_cm=3.2):
    with Image.open(image_path) as image:
        width, height = image.size
    if width <= 0 or height <= 0:
        raise ValueError("Ảnh không có kích thước hợp lệ")
    scale = min(max_width_cm / width, max_height_cm / height)
    return Cm(width * scale), Cm(height * scale)


def _fill_data_table(table, rows):
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    _set_row_repeat(table.rows[0])

    if not rows:
        row = table.add_row()
        row.cells[0].merge(row.cells[-1])
        paragraph = row.cells[0].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(
            paragraph.add_run("Không có trường hợp chấm công bất thường trong khoảng thời gian đã chọn."),
            size=10,
            italic=True,
        )
        _set_row_cant_split(row)
        return

    for row_index, item in enumerate(rows, 1):
        row = table.add_row()
        _set_row_cant_split(row)
        values = [item["name"], item["date"], item["explanation"], "", item["note"]]
        for cell_index, (cell, width) in enumerate(zip(row.cells, DATA_WIDTHS)):
            cell.width = width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell, top=90, bottom=90)
            if row_index % 2 == 0:
                _set_cell_shading(cell, "F3F7FB")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if cell_index in (0, 1, 3) else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_after = Pt(0)
            if values[cell_index]:
                _set_font(paragraph.add_run(values[cell_index]), size=9)

        image_cell = row.cells[3]
        image_paragraph = image_cell.paragraphs[0]
        image_path = item.get("image_path")
        if image_path and os.path.isfile(image_path):
            try:
                width, height = _fit_image(image_path)
                image_paragraph.add_run().add_picture(image_path, width=width, height=height)
            except Exception:
                _set_font(image_paragraph.add_run("Không thể đọc ảnh"), size=8, italic=True)
        else:
            _set_font(image_paragraph.add_run("Không có ảnh phù hợp"), size=8, italic=True)


def _resolve_period(persons, from_date, to_date):
    all_dates = [
        parsed
        for person in persons
        for record in person.get("records", [])
        if (parsed := _parse_date(record.get("date"))) is not None
    ]
    start = _parse_date(from_date) or (min(all_dates) if all_dates else date.today())
    end = _parse_date(to_date) or (max(all_dates) if all_dates else start)
    if start > end:
        raise ValueError("Từ ngày không được sau Đến ngày")
    return start, end


def _safe_filename_part(value):
    normalized = unicodedata.normalize("NFD", str(value or ""))
    ascii_text = "".join(char for char in normalized if unicodedata.category(char) != "Mn")
    ascii_text = ascii_text.replace("đ", "d").replace("Đ", "D").upper()
    return re.sub(r"[^A-Z0-9]+", "_", ascii_text).strip("_") or "DU_AN"


def export_aggregate_report(
    persons: Iterable[Dict],
    output_dir: str,
    project_name: str,
    from_date=None,
    to_date=None,
    report_date=None,
    template_path: str = DEFAULT_TEMPLATE_PATH,
) -> str:
    """Xuất một DOCX giải trình tổng hợp, giữ nguyên các file cá nhân hiện có."""
    persons = list(persons)
    start, end = _resolve_period(persons, from_date, to_date)
    report_day = _parse_date(report_date) or date.today()
    if not os.path.isfile(template_path):
        raise FileNotFoundError(f"Không tìm thấy template giải trình: {template_path}")

    doc = Document(template_path)
    _replace_tokens(doc, {
        "{{PROJECT_NAME}}": str(project_name or "").strip() or "Dự án",
        "{{FROM_DATE}}": start.strftime("%d/%m/%Y"),
        "{{TO_DATE}}": end.strftime("%d/%m/%Y"),
        "{{REPORT_DATE_LONG}}": report_day.strftime("Ngày %d tháng %m năm %Y"),
    })
    rows = collect_issue_rows(persons, from_date=start, to_date=end)
    _fill_data_table(_find_data_table(doc), rows)

    os.makedirs(output_dir, exist_ok=True)
    filename = (
        f"GIAI_TRINH_{_safe_filename_part(project_name)}_"
        f"{start.strftime('%d-%m-%Y')}_{end.strftime('%d-%m-%Y')}.docx"
    )
    output_path = os.path.join(output_dir, filename)
    doc.save(output_path)
    return output_path
