import io
import json
import zipfile
from urllib.parse import unquote

from docx import Document
from flask import Flask
from PIL import Image
import pytest

from src.daily_photo_routes import register_daily_photo_routes


def make_client(root, resolver=None):
    app = Flask(__name__)
    app.testing = True
    register_daily_photo_routes(app, lambda: root, project_resolver=resolver)
    return app.test_client()


def add_photo(root, project, day, name, shift=None, size=(80, 140)):
    folder = root / project / day
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / name
    Image.new('RGB', size, name.encode('utf-8')[0]).save(path)
    if shift:
        path.with_suffix(path.suffix + '.json').write_text(
            json.dumps({'shift': shift, 'send_time': '09:05:59'}), encoding='utf-8')
    return path


def report_text(doc):
    texts = [p.text for p in doc.paragraphs]
    texts.extend(c.text for t in doc.tables for r in t.rows for c in r.cells)
    return '\n'.join(texts).replace('\u200b', '')


def test_export_matches_gallery_and_is_scoped_to_project_and_date(tmp_path):
    client = make_client(tmp_path)
    for name, shift in [('sáng.png', 'morning'), ('chiều.png', 'afternoon'),
                        ('ngoài.png', 'outside'), ('cũ.png', None), ('lỗi.png', None)]:
        p = add_photo(tmp_path, 'Dự án A', '2026-09-18', name, shift)
        if name == 'lỗi.png':
            p.with_suffix('.png.json').write_text('bad metadata', encoding='utf-8')
    add_photo(tmp_path, 'Dự án B', '2026-09-18', 'other-project.png', 'morning')
    add_photo(tmp_path, 'Dự án A', '2026-08-18', 'other-month.png', 'morning')
    query = '?project=Dự án A&date=2026-09-18'
    response = client.get('/api/photos/daily/export-word' + query)
    assert response.status_code == 200
    assert response.mimetype.endswith('wordprocessingml.document')
    assert 'Anh_Dự án A_2026-09-18.docx' in unquote(response.headers['Content-Disposition'])
    doc = Document(io.BytesIO(response.data))
    text = report_text(doc)
    assert 'Tổng số: 5 ảnh' in text
    assert 'Ca sáng · 05:00–15:00 (1 ảnh)' in text
    assert 'Ca chiều · 16:00–24:00 (1 ảnh)' in text
    assert 'Ảnh khác (3 ảnh)' in text
    assert 'Giờ gửi:' not in text
    listing = client.get('/api/photos/daily/2026-09-18?project=Dự án A').json
    for photo in listing['photos']:
        assert photo['filename'] not in text
    assert 'other-project' not in text and 'other-month' not in text
    assert len(doc.inline_shapes) == listing['total']


@pytest.mark.parametrize('shift', ['morning', 'afternoon', 'outside', None])
def test_single_group_and_empty_shifts(tmp_path, shift):
    add_photo(tmp_path, 'Site', '2026-09-18', 'a.png', shift)
    response = make_client(tmp_path).get('/api/photos/daily/export-word?project=Site&date=2026-09-18')
    doc = Document(io.BytesIO(response.data))
    text = report_text(doc)
    assert 'Ca sáng' in text and 'Ca chiều' in text
    assert text.count('Không có ảnh') == (1 if shift in ('morning', 'afternoon') else 2)
    assert ('Ảnh khác' in text) == (shift not in ('morning', 'afternoon'))


def test_project_id_uses_display_name_and_mapped_legacy_day(tmp_path):
    add_photo(tmp_path, 'storage', '18', 'a.png', 'morning')
    (tmp_path / 'storage' / 'attendance-period.json').write_text(
        json.dumps({'version': 1, 'mappings': {'18': {
            'date': '2026-09-18', 'confirmed_by': 'test', 'confirmed_at': '2026-09-18T00:00:00Z'}}}))
    resolver = lambda ref: {'storage_dir': 'storage', 'display_name': 'Tên mới', 'active': True}
    client = make_client(tmp_path, resolver)
    response = client.get('/api/photos/daily/export-word?project_id=pid&date=2026-09-18')
    assert response.status_code == 200
    assert 'Tên mới' in report_text(Document(io.BytesIO(response.data)))
    assert 'Anh_Tên mới_2026-09-18.docx' in unquote(response.headers['Content-Disposition'])
    assert client.get('/api/photos/daily/export-word?project_id=pid&date=2026-08-18').status_code == 400


def test_empty_invalid_ambiguous_and_broken_images(tmp_path):
    client = make_client(tmp_path)
    for query in ['project=Site&date=2026-09-18', 'project=Site&date=18',
                  'project=Site&date=2026-02-30', 'project=../bad&date=2026-09-18']:
        assert client.get('/api/photos/daily/export-word?' + query).status_code == 400
    path = add_photo(tmp_path, 'Site', '18', 'legacy.png')
    assert client.get('/api/photos/daily/export-word?project=Site&date=2026-09-18').status_code == 409
    path = add_photo(tmp_path, 'Site', '2026-09-19', 'broken.png', 'morning')
    path.write_bytes(b'broken image')
    response = client.get('/api/photos/daily/export-word?project=Site&date=2026-09-19')
    assert response.status_code == 400
    assert 'broken.png' in response.json['error']


def test_many_photos_grid_order_ratio_and_row_pagination(tmp_path):
    for i in range(31):
        add_photo(tmp_path, 'Site', '2026-09-18', f'{i:02}.png', 'morning',
                  size=(160, 80) if i % 2 else (80, 160))
    add_photo(tmp_path, 'Site', '2026-09-18', 'pm.webp', 'afternoon')
    response = make_client(tmp_path).get('/api/photos/daily/export-word?project=Site&date=2026-09-18')
    doc = Document(io.BytesIO(response.data))
    assert len(doc.inline_shapes) == 32
    assert [len(t.rows) for t in doc.tables] == [4, 3, 1]
    assert all(len(r.cells) == 5 for r in doc.tables[0].rows)
    assert all(c.text == '' for t in doc.tables for r in t.rows for c in r.cells)
    assert all(len(c.paragraphs) == 1 for t in doc.tables for r in t.rows for c in r.cells)
    assert all(not p.paragraph_format.keep_with_next
               for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs)
    assert doc.sections[0].page_width.cm == pytest.approx(21, abs=.01)
    assert doc.sections[0].page_height.cm == pytest.approx(29.7, abs=.01)
    for i, shape in enumerate(list(doc.inline_shapes)[:31]):
        assert shape.width / shape.height == pytest.approx(2 if i % 2 else .5, rel=.01)
    headings = [p for p in doc.paragraphs if p.style.name == 'Heading 1']
    assert not headings[0].paragraph_format.page_break_before
    assert headings[1].paragraph_format.page_break_before
    assert 'tiếp theo' in headings[1].text
    assert headings[2].paragraph_format.page_break_before
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        xml = archive.read('word/document.xml').decode()
    assert xml.count('<w:cantSplit') == 8
    assert '<w:keepNext w:val="0"' in xml
    assert '<w:left w:w="40" w:type="dxa"' in xml
    assert doc.inline_shapes[1].width.cm == pytest.approx(3.45, abs=.01)
