from datetime import date, time

import pytest
from docx import Document
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook

from src.attendance_highlight import should_highlight_attendance
from src.excel_splitter import ExcelAttendanceSplitter
from src.excel_list_word_exporter import ExcelListWordExporter

HEADERS = ['Mã nhân viên', 'Tên nhân viên', 'Ngày', 'Giờ vào', 'Giờ ra']
COLUMNS = {'date': 2, 'gio_vao': 3, 'gio_ra': 4}
CASES = [
    ('', '17:00', True), ('07:00', '', True), ('', '', True),
    ('17:41', '17:41', True), ('07:00', '09:59', True),
    ('07:00', '10:00', True), ('07:00', '10:01', False),
    ('17:41', '06:00', False), ('23:00', '01:00', True),
    (time(7), time(10), True), (time(7), time(17), False),
    (time(7), time(10, 0, 1), False),
    (7 / 24, 10 / 24, True), (7 / 24, 17 / 24, False),
    ('07:00:00', '10:00:01', False), ('invalid', '17:00', False),
]


@pytest.mark.parametrize('start,end,expected', CASES)
def test_highlight_rule(start, end, expected):
    assert should_highlight_attendance(['1', 'Person', date(2026, 9, 1), start, end], COLUMNS) is expected


def test_non_data_rows_are_not_highlighted():
    for day in ('', 'Tổng cộng', 'Ngày'):
        assert not should_highlight_attendance(['1', 'Person', day, '', ''], COLUMNS)
    assert not should_highlight_attendance(['1', 'Person', '9/1/2026'], {'date': 2})


def test_excel_and_word_have_matching_full_row_highlights(tmp_path):
    source = tmp_path / 'source.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['CHI TIẾT CHẤM CÔNG'])
    sheet.append(HEADERS)
    for index, (start, end, _) in enumerate(CASES, 1):
        sheet.append(['1', 'Person', date(2026, 9, index), start, end])
    sheet.append(['1', 'Person', 'Tổng cộng', '', ''])
    workbook.save(source)
    original = source.read_bytes()

    files, _ = ExcelAttendanceSplitter(str(source)).split(str(tmp_path / 'split'))
    excel = load_workbook(files[0]).active
    word_path = ExcelListWordExporter(str(tmp_path / 'word')).export_from_excel(files[0])
    table = Document(word_path).tables[0]
    assert len(table.rows) == len(CASES) + 2
    assert source.read_bytes() == original
    for index, (_, _, expected) in enumerate(CASES):
        for cell in excel[index + 3]:
            assert (cell.fill.fgColor.rgb == '00FFF2CC') is expected
        for cell in table.rows[index + 1].cells:
            shade = cell._tc.get_or_add_tcPr().find(qn('w:shd'))
            assert (shade is not None and shade.get(qn('w:fill')) == 'FFF2CC') is expected
    assert all(cell.fill.patternType is None for cell in excel[2])
    assert all(cell.fill.patternType is None for cell in excel[excel.max_row])
    assert all(cell._tc.get_or_add_tcPr().find(qn('w:shd')) is None for cell in table.rows[-1].cells)
