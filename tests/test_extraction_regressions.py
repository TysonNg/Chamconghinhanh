from datetime import date, timedelta
from pathlib import Path
from types import SimpleNamespace

import pytest
import xlrd
from docx import Document
from openpyxl import Workbook, load_workbook

from src.excel_splitter import ExcelAttendanceSplitter, _ExcelReader
from src.excel_list_word_exporter import ExcelListWordExporter, _format_xlrd_cell
from src import pdf_extractor


@pytest.mark.parametrize("names", [("Alice", ""), ("", "Alice"), ("", "")])
def test_split_keeps_rows_with_code_and_blank_name(tmp_path, names):
    wb = Workbook()
    wb.active.append(["Mã nhân viên", "Tên nhân viên", "Ngày", "Giờ vào", "Giờ ra"])
    for day, name in enumerate(names, 1):
        wb.active.append(["001", name, f"{day:02d}/09/2026", "08:00", "17:00"])
    source = tmp_path / "input.xlsx"
    wb.save(source)
    paths, summaries = ExcelAttendanceSplitter(str(source)).split(str(tmp_path / "split"))
    assert len(paths) == 1
    assert summaries[0]["rows"] == 2
    assert summaries[0]["name"] == ("Alice" if "Alice" in names else "001")
    result = load_workbook(paths[0])
    assert result.active.max_row == 3
    assert [result.active.cell(r, 1).value for r in (2, 3)] == ["001", "001"]
    result.close()


@pytest.mark.parametrize("headers", [
    ["ID", "Họ và tên", "Ngày", "Check in", "Check out"],
    ["Mã thẻ", "Tên", "Ngày", "Giờ vào", "Giờ ra"],
    ["Tên nhân viên", "Ngày", "Giờ vào", "Giờ ra"],
])
def test_supported_split_headers_export_word(tmp_path, headers):
    wb = Workbook()
    wb.active.append(headers)
    wb.active.append((["001"] if len(headers) == 5 else []) + ["Alice", "01/09/2026", "08:00", "17:00"])
    source = tmp_path / "input.xlsx"
    wb.save(source)
    files, _ = ExcelAttendanceSplitter(str(source)).split(str(tmp_path / "split"))
    output = ExcelListWordExporter(str(tmp_path / "word")).export_from_excel(files[0])
    assert output is not None
    assert [c.text for c in Document(output).tables[0].rows[0].cells] == headers


def duration_workbook(fmt):
    return SimpleNamespace(datemode=0, xf_list=[SimpleNamespace(format_key=164)],
                           format_map={164: SimpleNamespace(format_str=fmt)})


@pytest.mark.parametrize("days", [1.5, 100.5])
def test_xls_elapsed_time_preserves_hours_through_split_and_word(tmp_path, days):
    cell = SimpleNamespace(ctype=xlrd.XL_CELL_DATE, value=days, xf_index=0)
    value = _ExcelReader._convert_xls_cell(None, duration_workbook("[h]:mm"), cell)
    assert value == timedelta(days=days)
    wb = Workbook()
    wb.active.append(["Mã nhân viên", "Tên nhân viên", "Ngày", "Tổng giờ"])
    wb.active.append(["001", "Alice", date(2026, 9, 1), value])
    path = tmp_path / "duration.xlsx"
    wb.save(path)
    output = ExcelListWordExporter(str(tmp_path / "word")).export_from_excel(str(path))
    assert Document(output).tables[0].rows[1].cells[3].text == f"{int(days * 24)}:00"


def test_direct_xls_word_elapsed_time_uses_cell_format():
    cell = SimpleNamespace(ctype=xlrd.XL_CELL_DATE, value=1.5, xf_index=0)
    assert _format_xlrd_cell(cell, 0, duration_workbook("[h]:mm")) == "36:00"


class InlineThread:
    def __init__(self, target, **kwargs):
        self.target = target
    def start(self):
        self.target()


def test_repeated_pdf_tasks_use_separate_batches(tmp_path, monkeypatch):
    batches = []
    def extract(source, output_dir, task):
        batches.append(Path(output_dir))
        Path(output_dir).mkdir(parents=True)
        (Path(output_dir) / "Alice.docx").write_bytes(b"generated")
    monkeypatch.setattr(pdf_extractor.threading, "Thread", InlineThread)
    monkeypatch.setattr(pdf_extractor, "extract_pdf_to_word", extract)
    root = tmp_path / "monthly"
    one = pdf_extractor.start_extraction_task("source.pdf", str(root))
    two = pdf_extractor.start_extraction_task("source.pdf", str(root))
    assert batches[0] != batches[1]
    assert all(len(list(batch.glob("*.docx"))) == 1 for batch in batches)
    assert pdf_extractor.get_task(one).output_dir == str(batches[0])
    assert pdf_extractor.get_task(two).output_dir == str(batches[1])


@pytest.mark.parametrize("empty", [False, True])
def test_excel_task_does_not_report_success_without_word(tmp_path, monkeypatch, empty):
    import src.app as app_module
    upload = tmp_path / "uploads"
    upload.mkdir()
    (upload / "input.xlsx").write_bytes(b"stub")
    monkeypatch.setattr(app_module, "EXCEL_UPLOAD_DIR", str(upload))
    monkeypatch.setattr(app_module, "EXCEL_PERSON_DIR", str(tmp_path / "persons"))
    monkeypatch.setattr(app_module, "EXCEL_OUTPUT_DIR", str(tmp_path / "word"))
    monkeypatch.setattr(app_module, "send_log", lambda *args: None)
    monkeypatch.setattr(app_module.threading, "Thread", InlineThread)
    monkeypatch.setattr(ExcelAttendanceSplitter, "__init__", lambda *args: None)
    monkeypatch.setattr(ExcelAttendanceSplitter, "split", lambda *args: ([], []) if empty else
                        (["person.xlsx"], [{"name": "Alice", "rows": 1, "present_rows": 1}]))
    monkeypatch.setattr(ExcelListWordExporter, "export_from_excel", lambda *args: None)
    response = app_module.app.test_client().post("/api/excel/extract", json={"filename": "input.xlsx"})
    assert response.status_code == 200
    result = app_module.app.test_client().get("/api/excel/status/" + response.json["task_id"]).json
    assert result["status"] == "failed"
    assert result["errors"]


def test_xls_date_with_time_keeps_display_and_epoch():
    from datetime import datetime
    from xlrd.xldate import xldate_from_datetime_tuple
    workbook = duration_workbook("m/d/yyyy hh:mm")
    workbook.datemode = 1
    serial = xldate_from_datetime_tuple((2026, 9, 1, 8, 30, 0), 1)
    cell = SimpleNamespace(ctype=xlrd.XL_CELL_DATE, value=serial, xf_index=0)
    assert _ExcelReader._convert_xls_cell(None, workbook, cell) == datetime(2026, 9, 1, 8, 30)
    assert _format_xlrd_cell(cell, 1, workbook) == "9/1/2026 08:30"


def test_word_uses_highest_scoring_header_like_splitter(tmp_path):
    wb = Workbook()
    wb.active.append(["Tên nhân viên", "Ngày"])
    wb.active.append(["ID", "Họ và tên", "Ngày", "Check in", "Check out"])
    wb.active.append(["001", "Alice", "01/09/2026", "08:00", "17:00"])
    source = tmp_path / "input.xlsx"
    wb.save(source)
    files, _ = ExcelAttendanceSplitter(str(source)).split(str(tmp_path / "split"))
    # Word can also export the original supported workbook.
    output = ExcelListWordExporter(str(tmp_path / "word")).export_from_excel(str(source))
    assert Document(output).tables[0].rows[0].cells[0].text == "ID"


def test_pdf_route_returns_actual_batch_and_keeps_same_name_pages(tmp_path, monkeypatch):
    import src.app as app_module
    upload = tmp_path / "uploads"
    upload.mkdir()
    (upload / "source.pdf").write_bytes(b"stub")
    monkeypatch.setattr(app_module, "PDF_UPLOAD_DIR", str(upload))
    monkeypatch.setattr(app_module, "PDF_OUTPUT_DIR", str(tmp_path / "outputs"))
    monkeypatch.setattr(app_module, "PDF_EXTRACTOR_AVAILABLE", True)
    monkeypatch.setattr(pdf_extractor, "PDF2DOCX_AVAILABLE", True)
    monkeypatch.setattr(pdf_extractor.threading, "Thread", InlineThread)
    class PDF:
        def __len__(self): return 2
        def close(self): pass
    class Converter:
        def __init__(self, path): pass
        def convert(self, path, start, end):
            document = Document()
            document.add_paragraph("Tên nhân viên: Alice Phòng ban: A")
            document.add_paragraph(f"page {start + 1}")
            document.save(path)
        def close(self): pass
    monkeypatch.setattr(pdf_extractor.fitz, "open", lambda path: PDF())
    monkeypatch.setattr(pdf_extractor, "Converter", Converter)
    client = app_module.app.test_client()
    batches = []
    for _ in range(2):
        response = client.post("/api/pdf/extract", json={"filename": "source.pdf"})
        assert response.status_code == 200
        task = pdf_extractor.get_task(response.json["task_id"])
        batch = Path(response.json["output_dir"])
        assert batch == Path(task.output_dir)
        assert task.status == "completed"
        assert len(list(batch.glob("*.docx"))) == 2
        assert {Document(f["path"]).paragraphs[1].text for f in task.files_created} == {"page 1", "page 2"}
        batches.append(batch)
    assert batches[0] != batches[1]
