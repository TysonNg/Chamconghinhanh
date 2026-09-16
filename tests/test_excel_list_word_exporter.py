from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from openpyxl import Workbook

from src.excel_list_word_exporter import ExcelListWordExporter


HEADERS = [
    "STT",
    "Mã nhân viên",
    "Tên nhân viên",
    "Phòng ban",
    "Ngày",
    "Thứ",
    "Giờ vào",
    "Giờ ra",
    "Trễ",
    "Sớm",
    "Công",
    "Tổng giờ",
    "Tăng ca",
    "Tổng toàn bộ",
    "Ca",
]


def _make_monthly_attendance(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["CHI TIẾT CHẤM CÔNG"])
    sheet.append(HEADERS)
    for day in range(1, 32):
        sheet.append([
            3479 + day,
            "2024063",
            "Bui thi bich nhung",
            "",
            f"8/{day}/2026",
            "Hai",
            "06:35",
            "16:51",
            0,
            9,
            1,
            8.8,
            0,
            10.3,
            "HChinh",
        ])
    workbook.save(path)


def test_monthly_word_table_uses_compact_fixed_widths(tmp_path):
    source = tmp_path / "attendance.xlsx"
    _make_monthly_attendance(source)

    output = ExcelListWordExporter(str(tmp_path)).export_from_excel(str(source))
    document = Document(output)
    section = document.sections[0]
    table = document.tables[0]

    assert round(section.page_width.cm, 1) == 29.7
    assert round(section.page_height.cm, 1) == 21.0
    assert len(table.columns) == 15
    assert len(table.rows) == 32
    assert table.autofit is False
    assert table._tbl.tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"

    grid_widths = [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid.gridCol_lst]
    printable_width = section.page_width.twips - section.left_margin.twips - section.right_margin.twips
    expected_widths = [
        Cm(value).twips
        for value in (1.3, 2.3, 4.4, 2.2, 2.0, 1.3, 1.6, 1.6, 1.0, 1.0, 1.0, 1.8, 1.7, 2.2, 1.6)
    ]
    assert grid_widths == expected_widths
    assert sum(grid_widths) <= printable_width

    assert [paragraph.text for paragraph in document.paragraphs] == ["CHI TIẾT CHẤM CÔNG"]
