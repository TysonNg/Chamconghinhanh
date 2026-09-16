# -*- coding: utf-8 -*-

import os
from datetime import date

from docx import Document
from docx.oxml.ns import qn
from openpyxl import Workbook
from PIL import Image

from src.aggregate_report_exporter import collect_issue_rows, export_aggregate_report
from src.excel_face_analyzer import ExcelFaceAnalyzer
from src.excel_extractor import ExcelToWordExporter


def test_collect_issue_rows_filters_normal_and_out_of_range_records():
    persons = [
        {
            "name": "Nguyễn Văn A",
            "records": [
                {
                    "date": "01/07/2026",
                    "gio_vao": "07:00",
                    "gio_ra": "19:00",
                    "is_absent": False,
                    "missing_checkin": False,
                    "missing_checkout": False,
                    "same_in_out": False,
                },
                {
                    "date": "05/07/2026",
                    "gio_vao": "",
                    "gio_ra": "19:03",
                    "is_absent": False,
                    "missing_checkin": True,
                    "missing_checkout": False,
                    "same_in_out": False,
                    "matched_image_path": "camera-05.jpg",
                },
                {
                    "date": "20/07/2026",
                    "gio_vao": "07:02",
                    "gio_ra": "07:04",
                    "is_absent": False,
                    "missing_checkin": False,
                    "missing_checkout": False,
                    "same_in_out": True,
                },
            ],
        }
    ]

    rows = collect_issue_rows(
        persons,
        from_date="2026-07-02",
        to_date="2026-07-10",
    )

    assert rows == [
        {
            "name": "Nguyễn Văn A",
            "date": "05/07/2026",
            "date_value": date(2026, 7, 5),
            "explanation": "Thiếu giờ vào, bổ sung hình ảnh thực tế",
            "note": "Giờ ra: 19:03",
            "image_path": "camera-05.jpg",
        }
    ]


def test_collect_issue_rows_labels_absence_as_absence():
    rows = collect_issue_rows([
        {
            "name": "Nguyễn Văn A",
            "records": [{
                "date": "07/07/2026",
                "gio_vao": "",
                "gio_ra": "",
                "is_absent": True,
                "missing_checkin": False,
                "missing_checkout": False,
                "same_in_out": False,
            }],
        }
    ])

    assert rows[0]["explanation"] == "Vắng mặt, cần giải trình"


def test_export_aggregate_report_creates_print_ready_a4_document(tmp_path):
    image_path = tmp_path / "camera.jpg"
    Image.new("RGB", (1200, 800), color=(45, 85, 125)).save(image_path)
    persons = [
        {
            "name": "Nguyễn Văn A",
            "records": [
                {
                    "date": "05/07/2026",
                    "gio_vao": "",
                    "gio_ra": "19:03",
                    "is_absent": False,
                    "missing_checkin": True,
                    "missing_checkout": False,
                    "same_in_out": False,
                    "matched_image_path": str(image_path),
                },
                {
                    "date": "06/07/2026",
                    "gio_vao": "07:01",
                    "gio_ra": "19:02",
                    "is_absent": False,
                    "missing_checkin": False,
                    "missing_checkout": False,
                    "same_in_out": False,
                },
            ],
        }
    ]

    output_path = export_aggregate_report(
        persons,
        output_dir=str(tmp_path),
        project_name="Chung cư Tân Thuận Đông",
        from_date="2026-07-01",
        to_date="2026-07-31",
        report_date=date(2026, 8, 5),
    )

    assert os.path.basename(output_path) == (
        "GIAI_TRINH_CHUNG_CU_TAN_THUAN_DONG_01-07-2026_31-07-2026.docx"
    )
    doc = Document(output_path)
    section = doc.sections[0]
    assert round(section.page_width.cm, 1) == 21.0
    assert round(section.page_height.cm, 1) == 29.7
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Chung cư Tân Thuận Đông" in body_text
    assert "Từ ngày 01/07/2026 đến ngày 31/07/2026" in body_text
    assert "Ngày 05 tháng 08 năm 2026" in "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )

    data_table = next(table for table in doc.tables if table.cell(0, 0).text == "TÊN")
    assert [cell.text for cell in data_table.rows[0].cells] == [
        "TÊN",
        "NGÀY",
        "GIẢI TRÌNH",
        "HÌNH ẢNH THỰC TẾ",
        "GHI CHÚ",
    ]
    assert len(data_table.rows) == 2
    assert data_table.cell(1, 0).text == "Nguyễn Văn A"
    assert data_table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
    assert data_table.rows[1]._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
    assert len(doc.inline_shapes) == 1
    signature_table = doc.tables[-1]
    for row in signature_table.rows[:3]:
        for cell in row.cells:
            assert cell.paragraphs[0]._p.get_or_add_pPr().find(qn("w:keepNext")) is not None


def test_person_export_records_the_camera_image_used_for_an_issue(tmp_path):
    portrait_dir = tmp_path / "portraits"
    image_dir = tmp_path / "camera" / "2026-07-05"
    output_dir = tmp_path / "output"
    portrait_dir.mkdir()
    image_dir.mkdir(parents=True)
    camera_path = image_dir / "matched.jpg"
    exif = Image.Exif()
    exif[36867] = "2026:07:05 08:00:00"
    Image.new("RGB", (640, 480), color=(20, 90, 45)).save(camera_path, exif=exif)
    from src.identity_registry import IdentityRegistry
    from src.face_matcher import MatchResult
    registry = IdentityRegistry(tmp_path / "identity.sqlite3", portrait_dir)
    project = registry.register_project("A")
    employee = registry.create_employee("Employee")
    registry.assign_employee(project["project_id"], employee["employee_id"], "001", "2026-01-01")

    class Matcher:
        def match_employee_in_images(self, **kwargs):
            assert kwargs["project_id"] == project["project_id"]
            assert kwargs["employee_id"] == employee["employee_id"]
            assert kwargs["camera_images"] == [str(camera_path)]
            return MatchResult(status="matched", image_path=str(camera_path), distance=0.1, project_id=project["project_id"], employee_id=employee["employee_id"])

    record = {
        "date": "05/07/2026",
        "weekday": "Chủ nhật",
        "gio_vao": "",
        "gio_ra": "19:03",
        "is_absent": False,
        "missing_checkin": True,
        "missing_checkout": False,
        "same_in_out": False,
    }
    exporter = ExcelToWordExporter(
        str(portrait_dir),
        str(output_dir),
        input_images_dir=str(tmp_path / "camera"),
        face_matcher=Matcher(), project_id=project["project_id"], identity_registry=registry,
    )

    exporter.export_person({"id": "001",
        "name": "Nguyễn Văn A",
        "month": 7,
        "year": 2026,
        "records": [record],
    })

    assert record["matched_image_path"] == str(camera_path)


def test_person_export_records_missing_camera_without_crashing(tmp_path):
    portrait_dir = tmp_path / "portraits"
    output_dir = tmp_path / "output"
    portrait_dir.mkdir()
    record = {
        "date": "08/07/2026",
        "weekday": "Thứ tư",
        "gio_vao": "07:10",
        "gio_ra": "",
        "is_absent": False,
        "missing_checkin": False,
        "missing_checkout": True,
        "same_in_out": False,
    }
    exporter = ExcelToWordExporter(
        str(portrait_dir),
        str(output_dir),
        input_images_dir=str(tmp_path / "camera-khong-ton-tai"),
        face_matcher=object(),
    )

    output_path = exporter.export_person({
        "name": "Nguyễn Văn B",
        "month": 7,
        "year": 2026,
        "records": [record],
    })

    assert os.path.isfile(output_path)
    assert record["matched_image_path"] is None


def test_excel_face_analyzer_adds_one_aggregate_report(tmp_path):
    input_dir = tmp_path / "excel"
    output_dir = tmp_path / "word"
    portrait_dir = tmp_path / "portraits"
    input_dir.mkdir()
    portrait_dir.mkdir()

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Mã nhân viên", "Tên nhân viên", "Ngày", "Thứ", "Giờ vào", "Giờ ra"])
    sheet.append(["NV01", "Nguyễn Văn A", "05/07/2026", "Chủ nhật", "", "19:03"])
    workbook.save(input_dir / "nguyen-van-a.xlsx")

    analyzer = ExcelFaceAnalyzer(
        str(portrait_dir),
        str(tmp_path / "camera"),
        matcher=None,
    )
    files = analyzer.analyze_folder(
        str(input_dir),
        str(output_dir),
        report_options={
            "project_name": "Chung cư Tân Thuận Đông",
            "from_date": "2026-07-01",
            "to_date": "2026-07-31",
            "report_date": date(2026, 8, 5),
        },
    )

    assert len(files) == 2
    aggregate = next(path for path in files if os.path.basename(path).startswith("GIAI_TRINH_"))
    assert os.path.isfile(aggregate)
    assert len(Document(aggregate).tables[0].rows) == 2


def test_export_many_rows_for_print_preview(tmp_path):
    landscape = tmp_path / "landscape.jpg"
    portrait = tmp_path / "portrait.jpg"
    Image.new("RGB", (1200, 800), color=(45, 85, 125)).save(landscape)
    Image.new("RGB", (800, 1200), color=(125, 75, 45)).save(portrait)
    records = []
    for day in range(1, 15):
        records.append({
            "date": f"{day:02d}/07/2026",
            "gio_vao": "" if day % 2 else "07:00",
            "gio_ra": "19:00" if day % 2 else "",
            "is_absent": False,
            "missing_checkin": bool(day % 2),
            "missing_checkout": not bool(day % 2),
            "same_in_out": False,
            "matched_image_path": (
                str(landscape) if day % 3 == 1
                else str(portrait) if day % 3 == 2
                else None
            ),
        })

    output_path = export_aggregate_report(
        [{"name": "Nguyễn Văn Nhân Viên Có Tên Dài", "records": records}],
        output_dir=str(tmp_path),
        project_name="Chung cư Tân Thuận Đông",
        from_date="2026-07-01",
        to_date="2026-07-31",
        report_date=date(2026, 8, 5),
    )

    doc = Document(output_path)
    assert len(doc.tables[0].rows) == 15
    assert len(doc.inline_shapes) == 10
